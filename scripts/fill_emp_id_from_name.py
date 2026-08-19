"""
Fill an Excel EMP_ID column by matching NAME against HRMS employees.

Expected NAME format: "LASTNAME, Firstname" (middle/suffix tolerated and ignored for matching).

This script is idempotent:
- If EMP_ID already has a value, it is left unchanged.
- Writes match status to a new column MATCH_STATUS (created if missing).

Usage (PowerShell):
  python scripts/fill_emp_id_from_name.py "C:\\python\\back-up\\hrms\\LEAVE CARD REGULAR.xlsx"
"""

from __future__ import annotations

import re
import sys
from pathlib import Path


def _norm_token(s: str) -> str:
    s = (s or "").strip().upper()
    s = s.replace("\u00a0", " ")
    s = re.sub(r"[.]", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _canon_header_key(s: str) -> str:
    """Uppercase and remove common separators for header matching."""
    t = _norm_token(s)
    t = t.replace(" ", "").replace("_", "").replace("-", "")
    return t


def _parse_name_cell(raw: str) -> tuple[str, str] | None:
    """
    Parse "LAST, FIRST ..." -> (LAST, FIRST)
    Returns None if not parseable.
    """
    text = (raw or "").strip()
    if not text:
        return None
    if "," not in text:
        return None
    last, rest = text.split(",", 1)
    last_n = _norm_token(last)
    rest = _norm_token(rest)
    if not last_n or not rest:
        return None
    # FIRST may include middle name/initials/suffix; take first token
    first_n = rest.split(" ", 1)[0].strip()
    if not first_n:
        return None
    return last_n, first_n


def _header_map(row_values: list[str | None]) -> dict[str, int]:
    m: dict[str, int] = {}
    for idx, v in enumerate(row_values):
        key = _norm_token(str(v or ""))
        if key:
            m[key] = idx
    return m


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print("Usage: python scripts/fill_emp_id_from_name.py <path_to_xlsx>")
        return 2

    xlsx_path = Path(argv[1]).expanduser()
    if not xlsx_path.exists():
        print(f"[ERROR] File not found: {xlsx_path}")
        return 2

    try:
        import openpyxl  # type: ignore
    except Exception as exc:
        print(f"[ERROR] openpyxl is required: {exc}")
        return 2

    # Load employee lookup from HRMS database
    try:
        # Ensure project root is importable when script is run directly.
        root = Path(__file__).resolve().parents[1]
        if str(root) not in sys.path:
            sys.path.insert(0, str(root))

        from app import create_app  # type: ignore
        from app.models import Employee  # type: ignore

        app = create_app()
        with app.app_context():
            employees = (
                Employee.query.with_entities(Employee.employee_id, Employee.last_name, Employee.first_name)
                .all()
            )
    except Exception as exc:
        print(f"[ERROR] Could not load employees from DB: {exc}")
        return 1

    lookup: dict[tuple[str, str], list[str]] = {}
    for emp_id, last, first in employees:
        key = (_norm_token(last or ""), _norm_token(first or ""))
        if not key[0] or not key[1]:
            continue
        lookup.setdefault(key, []).append(str(emp_id))

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    # Detect header row by scanning the top area (some files have title rows).
    header_row = None
    header = {}
    for r in range(1, min(ws.max_row, 40) + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        h = _header_map([str(v) if v is not None else "" for v in vals])
        # Find NAME and EMP_ID columns with flexible header spellings.
        name_col = None
        empid_col = None
        for k, idx in h.items():
            ck = _canon_header_key(k)
            if ck == "NAME":
                name_col = idx
            if ck in ("EMPID", "EMPLOYEEID"):
                empid_col = idx
        if name_col is not None and empid_col is not None:
            header_row = r
            header = h  # keep original tokens for MATCH_STATUS lookup
            break

    if header_row is None:
        # Show whatever is on row 1 to help debugging.
        vals0 = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        h0 = _header_map([str(v) if v is not None else "" for v in vals0])
        print("[ERROR] Could not find a header row containing NAME and EMP_ID.")
        print(f"Row 1 headers detected: {sorted(h0.keys())}")
        return 2

    # Find columns (case-insensitive)
    name_col = None
    empid_col = None
    for k, idx in header.items():
        ck = _canon_header_key(k)
        if ck == "NAME":
            name_col = idx
        if ck in ("EMPID", "EMPLOYEEID"):
            empid_col = idx
    status_col = header.get("MATCH_STATUS")

    # openpyxl columns are 1-based; our header map indices are 0-based
    name_col_1 = name_col + 1
    empid_col_1 = empid_col + 1

    if status_col is None:
        status_col_1 = ws.max_column + 1
        ws.cell(row=header_row, column=status_col_1).value = "MATCH_STATUS"
    else:
        status_col_1 = status_col + 1

    filled = 0
    skipped_existing = 0
    no_match = 0
    multiple = 0
    bad_name = 0

    for r in range(header_row + 1, ws.max_row + 1):
        name_val = ws.cell(row=r, column=name_col_1).value
        emp_val = ws.cell(row=r, column=empid_col_1).value

        if emp_val not in (None, "", 0, "0"):
            skipped_existing += 1
            continue

        parsed = _parse_name_cell(str(name_val or ""))
        if not parsed:
            bad_name += 1
            ws.cell(row=r, column=status_col_1).value = "BAD_NAME"
            continue

        key = parsed
        hits = lookup.get(key) or []
        if not hits:
            no_match += 1
            ws.cell(row=r, column=status_col_1).value = "NO_MATCH"
            continue
        if len(hits) > 1:
            multiple += 1
            ws.cell(row=r, column=status_col_1).value = f"MULTIPLE({len(hits)})"
            continue

        ws.cell(row=r, column=empid_col_1).value = hits[0]
        ws.cell(row=r, column=status_col_1).value = "OK"
        filled += 1

    out_path = xlsx_path.with_name(xlsx_path.stem + "_filled" + xlsx_path.suffix)
    wb.save(out_path)

    print(f"[OK] Saved: {out_path}")
    print(f"Filled: {filled}")
    print(f"Skipped (already had EMP_ID): {skipped_existing}")
    print(f"No match: {no_match}")
    print(f"Multiple matches: {multiple}")
    print(f"Bad/blank name: {bad_name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))

