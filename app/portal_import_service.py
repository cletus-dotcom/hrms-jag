"""Generic CSV/Excel import, export, and purge for HRMS database tables."""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Any

from flask import current_app
from sqlalchemy import Boolean, Date, DateTime, Integer, Numeric, Time, inspect
from sqlalchemy.exc import SQLAlchemyError
from werkzeug.datastructures import FileStorage

from app import db

IMPORT_SKIP_COLUMNS = {'password_hash'}
IMPORT_SENSITIVE_TABLES = {'users'}
PASSWORD_IMPORT_COLUMN = 'password'
IMPORT_RESOLVER_VERSION = 'employee-badge-v2'

# Columns named employee_id on other tables reference employees.id (PK), not employees.employee_id.
_EMPLOYEE_PK_CACHE: dict[str, int] | None = None
_USER_PK_CACHE: dict[str, int] | None = None


def _clear_fk_caches() -> None:
    global _EMPLOYEE_PK_CACHE, _USER_PK_CACHE
    _EMPLOYEE_PK_CACHE = None
    _USER_PK_CACHE = None


def _employee_pk_lookup() -> dict[str, int]:
    global _EMPLOYEE_PK_CACHE
    if _EMPLOYEE_PK_CACHE is None:
        from app.models import Employee

        lookup: dict[str, int] = {}
        for pk, badge in db.session.query(Employee.id, Employee.employee_id).all():
            lookup[str(pk)] = pk
            if badge is not None:
                lookup[str(badge).strip()] = pk
        _EMPLOYEE_PK_CACHE = lookup
    return _EMPLOYEE_PK_CACHE


def _user_pk_lookup() -> dict[str, int]:
    global _USER_PK_CACHE
    if _USER_PK_CACHE is None:
        from app.models import User

        lookup: dict[str, int] = {}
        for pk, badge, username in db.session.query(
            User.id, User.employee_id, User.username
        ).all():
            lookup[str(pk)] = pk
            if badge is not None:
                lookup[str(badge).strip()] = pk
            if username is not None:
                lookup[str(username).strip()] = pk
        _USER_PK_CACHE = lookup
    return _USER_PK_CACHE


def _resolve_import_fk(model, column_name: str, column, value: Any) -> Any:
    """Map human-facing IDs (employee badge, user login id) to FK primary keys."""
    if value is None or value == '':
        return None
    table_name = model.__tablename__
    if column_name == 'employee_id' and table_name != 'employees':
        key = str(value).strip()
        if key.endswith('.0') and key[:-2].isdigit():
            key = key[:-2]
        resolved = _employee_pk_lookup().get(key)
        if resolved is None:
            raise ValueError(
                f'No employee found for id/employee_id={value}. '
                'Import employees first or use the internal employees.id from export.'
            )
        return resolved
    if column_name in ('created_by', 'user_id', 'deleted_by_user_id') and table_name != 'users':
        key = str(value).strip()
        if key.endswith('.0') and key[:-2].isdigit():
            key = key[:-2]
        resolved = _user_pk_lookup().get(key)
        if resolved is None:
            raise ValueError(
                f'No user found for id/employee_id/username={value}.'
            )
        return resolved
    return value


def _apply_import_fk_resolution(model, mapped: dict[str, Any], errors: list[str]) -> None:
    col_by_name = {c.name: c for c in model.__table__.columns}
    for key, val in list(mapped.items()):
        col = col_by_name.get(key)
        if col is None or val is None:
            continue
        try:
            mapped[key] = _resolve_import_fk(model, key, col, val)
        except ValueError as exc:
            errors.append(str(exc))


def _iter_models():
    import app.models as models_pkg

    for name in sorted(dir(models_pkg)):
        obj = getattr(models_pkg, name)
        if (
            isinstance(obj, type)
            and hasattr(obj, '__tablename__')
            and getattr(obj, '__table__', None) is not None
        ):
            yield obj


def list_importable_tables():
    items = []
    for model in _iter_models():
        table = model.__table__
        cols = [c.name for c in table.columns if c.name not in IMPORT_SKIP_COLUMNS]
        items.append({
            'key': model.__tablename__,
            'label': model.__name__,
            'column_count': len(cols),
            'row_count': db.session.query(model).count(),
            'sensitive': model.__tablename__ in IMPORT_SENSITIVE_TABLES,
        })
    return items


def get_model_by_table(table_key: str):
    for model in _iter_models():
        if model.__tablename__ == table_key:
            return model
    return None


def importable_columns(model) -> list:
    extra = [PASSWORD_IMPORT_COLUMN] if model.__tablename__ == 'users' else []
    cols = [c.name for c in model.__table__.columns if c.name not in IMPORT_SKIP_COLUMNS]
    return cols + extra


def build_template_csv(model) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(importable_columns(model))
    writer.writerow([])
    return buf.getvalue()


def build_template_xlsx(model) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = (model.__tablename__)[:31]
    ws.append(importable_columns(model))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _normalize_header(h: str) -> str:
    return re.sub(r'[\s_]+', '', (h or '').strip().lower())


def _map_headers(model, raw_headers: list[str]) -> dict[str, str]:
    columns = {_normalize_header(c): c for c in importable_columns(model)}
    mapping = {}
    for h in raw_headers:
        key = _normalize_header(h)
        if key in columns:
            mapping[h] = columns[key]
    return mapping


def _cell_to_import_text(cell) -> str:
    """Normalize Excel/CSV cell values to strings import coercion understands."""
    if cell is None:
        return ''
    if isinstance(cell, datetime):
        if (
            cell.hour == 0 and cell.minute == 0 and cell.second == 0
            and cell.microsecond == 0
        ):
            return cell.strftime('%Y-%m-%d')
        if cell.microsecond:
            return cell.strftime('%Y-%m-%d %H:%M:%S.%f')
        return cell.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(cell, date):
        return cell.strftime('%Y-%m-%d')
    if isinstance(cell, time):
        return cell.strftime('%H:%M:%S')
    if isinstance(cell, float):
        if cell == int(cell):
            return str(int(cell))
    if isinstance(cell, bool):
        return '1' if cell else '0'
    return str(cell).strip()


def _json_safe(value: Any) -> Any:
    """Convert preview payload values to JSON-serializable forms."""
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=' ', timespec='seconds')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime('%H:%M:%S')
    if isinstance(value, dict):
        return {k: _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    return str(value)


def _coerce_value(column, raw: str):
    if raw is None:
        return None
    text = str(raw).strip()
    if text == '' or text.lower() in ('null', 'none'):
        return None
    col_type = type(column.type)
    if col_type is Boolean:
        return text.lower() in ('1', 'true', 'yes', 'y', 't')
    if col_type is Integer:
        return int(float(text))
    if col_type is Numeric:
        try:
            return Decimal(text)
        except InvalidOperation:
            raise ValueError(f'Invalid number: {text}')
    if col_type is Date:
        for fmt in (
            '%Y-%m-%d',
            '%m/%d/%Y',
            '%d/%m/%Y',
            '%Y-%m-%d %H:%M:%S',
            '%Y-%m-%d %H:%M:%S.%f',
        ):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        raise ValueError(f'Invalid date: {text}')
    if col_type is DateTime:
        for fmt in (
            '%Y-%m-%d %H:%M:%S.%f',
            '%Y-%m-%d %H:%M:%S',
            '%Y-%m-%dT%H:%M:%S',
            '%Y-%m-%d',
        ):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
        raise ValueError(f'Invalid datetime: {text}')
    if col_type is Time:
        for fmt in ('%H:%M:%S', '%H:%M'):
            try:
                return datetime.strptime(text, fmt).time()
            except ValueError:
                continue
        raise ValueError(f'Invalid time: {text}')
    return text


def _read_rows_from_upload(model, file_storage: FileStorage) -> tuple[list[str], list[dict[str, str]]]:
    filename = (file_storage.filename or '').lower()
    if hasattr(file_storage, 'stream') and hasattr(file_storage.stream, 'seek'):
        file_storage.stream.seek(0)
    if filename.endswith(('.xlsx', '.xlsm')):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError('Excel support requires openpyxl. Install it in the server environment.') from exc
        try:
            wb = load_workbook(file_storage, read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(f'Could not read Excel file: {exc}') from exc
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [_cell_to_import_text(c) for c in next(rows_iter)]
        except StopIteration:
            wb.close()
            return [], []
        data_rows = []
        for row in rows_iter:
            if not any(cell is not None and _cell_to_import_text(cell) for cell in row):
                continue
            data_rows.append({
                headers[i]: _cell_to_import_text(cell)
                for i, cell in enumerate(row)
                if i < len(headers)
            })
        wb.close()
        return headers, data_rows
    content = file_storage.read()
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = content.decode('latin-1')
    reader = csv.DictReader(io.StringIO(text))
    headers = reader.fieldnames or []
    return list(headers), [dict(row) for row in reader]


def preview_import(model, file_storage: FileStorage) -> dict:
    try:
        raw_headers, raw_rows = _read_rows_from_upload(model, file_storage)
    except ValueError as exc:
        return {'ok': False, 'error': str(exc)}
    except Exception as exc:
        current_app.logger.exception('Import file read failed for %s', model.__tablename__)
        return {'ok': False, 'error': f'Could not read upload file: {exc}'}
    header_map = _map_headers(model, raw_headers)
    if not header_map:
        return {'ok': False, 'error': 'No recognizable column headers found. Download the template first.'}

    _clear_fk_caches()
    table = model.__table__
    col_by_name = {c.name: c for c in table.columns}
    preview_rows = []
    valid_count = 0
    for idx, raw in enumerate(raw_rows[:500], start=1):
        mapped: dict[str, Any] = {}
        errors = []
        for src, dest in header_map.items():
            raw_val = raw.get(src, '')
            if dest == PASSWORD_IMPORT_COLUMN:
                mapped[dest] = str(raw_val).strip()
                continue
            col = col_by_name.get(dest)
            if col is None:
                continue
            try:
                mapped[dest] = _coerce_value(col, raw_val)
            except (ValueError, TypeError) as exc:
                errors.append(f'{dest}: {exc}')
        if model.__tablename__ == 'users' and not mapped.get(PASSWORD_IMPORT_COLUMN):
            errors.append('password: required for new user rows')
        _apply_import_fk_resolution(model, mapped, errors)
        preview_rows.append({
            'row_num': idx,
            'data': _json_safe(mapped),
            'errors': errors,
            'valid': len(errors) == 0,
        })
        if not errors:
            valid_count += 1

    return {
        'ok': True,
        'table': model.__tablename__,
        'import_resolver': IMPORT_RESOLVER_VERSION,
        'columns': list(header_map.values()),
        'rows': preview_rows,
        'total': len(raw_rows),
        'preview_count': len(preview_rows),
        'valid_count': valid_count,
        'invalid_count': len(preview_rows) - valid_count,
        'truncated': len(raw_rows) > len(preview_rows),
    }


def commit_import(model, rows: list[dict]) -> dict:
    if not rows:
        return {'ok': False, 'error': 'No rows to import.'}
    invalid = [r for r in rows if not r.get('valid', True)]
    if invalid:
        return {'ok': False, 'error': f'{len(invalid)} invalid row(s). Fix preview errors first.'}

    _clear_fk_caches()
    inserted = 0
    errors = []
    col_by_name = {c.name: c for c in model.__table__.columns}
    for item in rows:
        payload = dict(item.get('data') or {})
        plain_password = None
        if model.__tablename__ == 'users':
            plain_password = payload.pop(PASSWORD_IMPORT_COLUMN, None) or payload.pop('password', None)
        row_errors: list[str] = []
        for key, val in list(payload.items()):
            col = col_by_name.get(key)
            if col is None or val is None:
                continue
            try:
                payload[key] = _resolve_import_fk(model, key, col, val)
            except ValueError as exc:
                row_errors.append(str(exc))
        if row_errors:
            errors.append(f"Row {item.get('row_num', '?')}: {'; '.join(row_errors)}")
            if len(errors) >= 10:
                break
            continue
        try:
            obj = model()
            for key, val in payload.items():
                if hasattr(obj, key):
                    setattr(obj, key, val)
            if model.__tablename__ == 'users' and plain_password:
                obj.set_password(str(plain_password))
            db.session.add(obj)
            db.session.flush()
            inserted += 1
        except (SQLAlchemyError, ValueError, TypeError) as exc:
            db.session.rollback()
            errors.append(f"Row {item.get('row_num', '?')}: {exc}")
            if len(errors) >= 10:
                break
    if errors:
        db.session.rollback()
        return {'ok': False, 'error': '; '.join(errors)}
    try:
        db.session.commit()
    except SQLAlchemyError as exc:
        db.session.rollback()
        current_app.logger.exception('Import commit failed')
        return {'ok': False, 'error': str(exc)}
    return {'ok': True, 'inserted': inserted}


def export_table_csv(model, limit: int = 5000) -> str:
    cols = [c.name for c in model.__table__.columns if c.name not in IMPORT_SKIP_COLUMNS]
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(cols)
    for row in db.session.query(model).limit(limit).all():
        writer.writerow([getattr(row, c, '') for c in cols])
    return buf.getvalue()


def purge_table(model, confirm_text: str) -> dict:
    expected = f"PURGE {model.__tablename__.upper()}"
    if (confirm_text or '').strip() != expected:
        return {'ok': False, 'error': f'Type exactly: {expected}'}
    try:
        count = db.session.query(model).count()
        db.session.query(model).delete()
        db.session.commit()
        return {'ok': True, 'deleted': count, 'table': model.__tablename__}
    except SQLAlchemyError as exc:
        db.session.rollback()
        return {'ok': False, 'error': str(exc)}
