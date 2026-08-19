"""Parse and persist Flag Raising / Flag Retreat / Meeting attendance Excel uploads."""
from __future__ import annotations

import time
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

import requests
from flask_login import current_user

from app import db
from app.models import (
    EVENT_ATTENDANCE_CATEGORIES,
    EVENT_ATTENDANCE_CATEGORY_LABELS,
    Employee,
    EventAttendance,
    EventAttendanceBatch,
)

PH_TZ = timezone(timedelta(hours=8))
_NOMINATIM_URL = 'https://nominatim.openstreetmap.org/reverse'
_NOMINATIM_UA = 'HRMS-LGU-Jagna/1.0 (event-attendance reverse-geocode)'
# Round coords for cache (~11 m) so nearby punches share one lookup.
_GEO_CACHE_DECIMALS = 4


def category_label(category: str) -> str:
    return EVENT_ATTENDANCE_CATEGORY_LABELS.get(category, category or '—')


def normalize_category(raw: str | None) -> str | None:
    s = (raw or '').strip().lower().replace(' ', '_').replace('-', '_')
    aliases = {
        'flag_raising': 'flag_raising',
        'flagraising': 'flag_raising',
        'raising': 'flag_raising',
        'flag_retreat': 'flag_retreat',
        'flagretreat': 'flag_retreat',
        'retreat': 'flag_retreat',
        'meeting': 'meeting',
        'meetings': 'meeting',
    }
    key = aliases.get(s, s)
    return key if key in EVENT_ATTENDANCE_CATEGORIES else None


def _header_map(row: tuple) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(row or ()):
        key = str(cell or '').strip().lower().replace(' ', '')
        if not key:
            continue
        mapping[key] = idx
    return mapping


def _cell(row: tuple, headers: dict[str, int], *names: str):
    for name in names:
        idx = headers.get(name)
        if idx is not None and idx < len(row):
            return row[idx]
    return None


def _parse_timestamp_ms(raw) -> int | None:
    if raw is None or raw == '':
        return None
    if isinstance(raw, datetime):
        return int(raw.replace(tzinfo=PH_TZ).timestamp() * 1000)
    try:
        val = Decimal(str(raw).strip())
    except (InvalidOperation, ValueError):
        return None
    # Epoch ms (13 digits) or seconds (10 digits)
    if val > Decimal('1e11'):
        return int(val)
    if val > Decimal('1e9'):
        return int(val * 1000)
    return None


def _ms_to_ph_datetime(timestamp_ms: int) -> datetime:
    return datetime.fromtimestamp(timestamp_ms / 1000.0, tz=PH_TZ).replace(tzinfo=None)


def _parse_coord(raw) -> Decimal | None:
    if raw is None or raw == '':
        return None
    try:
        return Decimal(str(raw).strip())
    except (InvalidOperation, ValueError):
        return None


def _badge_lookup_keys(badge: str) -> set[str]:
    badge = (badge or '').strip()
    keys = {badge} if badge else set()
    if badge.isdigit():
        n = int(badge)
        keys.add(str(n))
        keys.add(f'{n:05d}')
        keys.add(f'{n:06d}')
    return {k for k in keys if k}


def build_badge_employee_map(badges: set[str]) -> dict[str, Employee]:
    """Map any badge variant from the upload to an Employee row."""
    lookup_keys: set[str] = set()
    for b in badges:
        lookup_keys |= _badge_lookup_keys(b)
    if not lookup_keys:
        return {}
    employees = Employee.query.filter(Employee.employee_id.in_(list(lookup_keys))).all()
    by_badge: dict[str, Employee] = {}
    for emp in employees:
        for key in _badge_lookup_keys(emp.employee_id or ''):
            by_badge[key] = emp
    return by_badge


def _coord_cache_key(lat, lon) -> tuple[float, float] | None:
    try:
        return (
            round(float(lat), _GEO_CACHE_DECIMALS),
            round(float(lon), _GEO_CACHE_DECIMALS),
        )
    except (TypeError, ValueError):
        return None


def _pick_address_part(addr: dict, *keys: str) -> str:
    for key in keys:
        val = (addr.get(key) or '').strip()
        if val:
            return val
    return ''


def format_ph_location(addr: dict | None) -> dict:
    """
    Build barangay / municipality / province from Nominatim address details.
    Returns dict with barangay, municipality, province, location_text.
    """
    addr = addr or {}
    barangay = _pick_address_part(
        addr, 'quarter', 'village', 'suburb', 'neighbourhood', 'neighborhood', 'hamlet', 'barangay'
    )
    municipality = _pick_address_part(
        addr, 'town', 'city', 'municipality', 'city_district', 'county'
    )
    province = _pick_address_part(addr, 'state', 'province', 'region')
    parts = []
    for part in (barangay, municipality, province):
        if part and part not in parts:
            parts.append(part)
    return {
        'barangay': barangay or None,
        'municipality': municipality or None,
        'province': province or None,
        'location_text': ', '.join(parts) if parts else None,
    }


def reverse_geocode_location(lat, lon, *, cache: dict | None = None, sleep_s: float = 1.05) -> dict:
    """
    Reverse-geocode lat/lon to barangay, town, province (OpenStreetMap Nominatim).
    Uses an in-memory cache keyed by rounded coordinates. Respects ~1 req/sec.
    """
    empty = {'barangay': None, 'municipality': None, 'province': None, 'location_text': None}
    key = _coord_cache_key(lat, lon)
    if key is None:
        return empty
    if cache is not None and key in cache:
        return cache[key]

    # Prefer previously stored location for the same rounded coords
    existing = (
        EventAttendance.query
        .filter(EventAttendance.location_text.isnot(None))
        .filter(EventAttendance.latitude.isnot(None))
        .filter(EventAttendance.longitude.isnot(None))
        .order_by(EventAttendance.id.desc())
        .limit(200)
        .all()
    )
    for row in existing:
        if _coord_cache_key(row.latitude, row.longitude) == key and row.location_text:
            result = {
                'barangay': row.barangay,
                'municipality': row.municipality,
                'province': row.province,
                'location_text': row.location_text,
            }
            if cache is not None:
                cache[key] = result
            return result

    try:
        if sleep_s and cache is not None and any(cache.values()):
            time.sleep(sleep_s)
        resp = requests.get(
            _NOMINATIM_URL,
            params={
                'lat': key[0],
                'lon': key[1],
                'format': 'json',
                'addressdetails': 1,
                'zoom': 18,
            },
            headers={'User-Agent': _NOMINATIM_UA},
            timeout=20,
        )
        if resp.status_code != 200:
            result = empty
        else:
            result = format_ph_location((resp.json() or {}).get('address'))
    except Exception:
        result = empty

    if cache is not None:
        cache[key] = result
    return result


def fill_missing_locations(records: list[EventAttendance], *, max_lookups: int = 40) -> int:
    """Backfill location_text for records missing it. Returns number updated."""
    cache: dict = {}
    updated = 0
    lookups = 0
    for rec in records:
        if (rec.location_text or '').strip():
            continue
        if rec.latitude is None or rec.longitude is None:
            continue
        if lookups >= max_lookups:
            break
        loc = reverse_geocode_location(rec.latitude, rec.longitude, cache=cache)
        lookups += 1
        if not loc.get('location_text'):
            continue
        rec.barangay = loc.get('barangay')
        rec.municipality = loc.get('municipality')
        rec.province = loc.get('province')
        rec.location_text = loc.get('location_text')
        updated += 1
    if updated:
        db.session.commit()
    return updated


def parse_attendance_excel(file_storage) -> list[dict]:
    """
    Read Excel with headers: employeeId, timestamp, lat, long.
    Returns list of row dicts (badge, timestamp_ms, punched_at, lat, long).
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError('Excel support requires openpyxl.') from exc

    raw = file_storage.read()
    if hasattr(file_storage, 'seek'):
        try:
            file_storage.seek(0)
        except Exception:
            pass
    wb = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError('Excel file is empty.') from exc

    headers = _header_map(header_row)
    if 'employeeid' not in headers or 'timestamp' not in headers:
        raise ValueError(
            'Excel must include columns employeeId and timestamp '
            '(optional: lat, long).'
        )

    parsed: list[dict] = []
    for row in rows_iter:
        if not row or all(c is None or str(c).strip() == '' for c in row):
            continue
        badge_raw = _cell(row, headers, 'employeeid', 'employee_id', 'empid', 'id')
        badge = str(badge_raw).strip() if badge_raw is not None else ''
        if badge.endswith('.0') and badge[:-2].isdigit():
            badge = badge[:-2]
        ts_ms = _parse_timestamp_ms(_cell(row, headers, 'timestamp', 'time', 'ts'))
        if not badge or ts_ms is None:
            continue
        parsed.append({
            'badge_employee_id': badge,
            'timestamp_ms': ts_ms,
            'punched_at': _ms_to_ph_datetime(ts_ms),
            'latitude': _parse_coord(_cell(row, headers, 'lat', 'latitude')),
            'longitude': _parse_coord(_cell(row, headers, 'long', 'lng', 'longitude', 'lon')),
        })
    if not parsed:
        raise ValueError('No valid attendance rows found in the Excel file.')
    return parsed


def upload_event_attendance(
    *,
    file_storage,
    category: str,
    event_date: date,
    filename: str = '',
) -> EventAttendanceBatch:
    cat = normalize_category(category)
    if not cat:
        raise ValueError(
            'Category must be Flag Raising, Flag Retreat, or Meeting.'
        )
    if not event_date:
        raise ValueError('Event date is required.')

    rows = parse_attendance_excel(file_storage)
    badges = {r['badge_employee_id'] for r in rows}
    emp_map = build_badge_employee_map(badges)

    existing_keys = {
        (r.badge_employee_id, r.timestamp_ms)
        for r in (
            EventAttendance.query
            .filter_by(category=cat, event_date=event_date)
            .filter(EventAttendance.badge_employee_id.in_(list(badges)))
            .all()
        )
    }

    user_id = getattr(current_user, 'id', None) if current_user and getattr(current_user, 'is_authenticated', False) else None
    batch = EventAttendanceBatch(
        category=cat,
        event_date=event_date,
        filename=(filename or getattr(file_storage, 'filename', '') or '')[:255],
        row_count=len(rows),
        uploaded_by_user_id=user_id if isinstance(user_id, int) else None,
    )
    db.session.add(batch)
    db.session.flush()

    inserted = 0
    skipped = 0
    unmatched = 0
    seen_in_file: set[tuple[str, int]] = set()
    geo_cache: dict = {}

    for row in rows:
        badge = row['badge_employee_id']
        ts = row['timestamp_ms']
        key = (badge, ts)
        if key in seen_in_file or key in existing_keys:
            skipped += 1
            continue
        seen_in_file.add(key)
        existing_keys.add(key)

        emp = None
        for k in _badge_lookup_keys(badge):
            emp = emp_map.get(k)
            if emp:
                break
        if not emp:
            unmatched += 1

        loc = {'barangay': None, 'municipality': None, 'province': None, 'location_text': None}
        if row['latitude'] is not None and row['longitude'] is not None:
            loc = reverse_geocode_location(
                row['latitude'], row['longitude'], cache=geo_cache
            )

        db.session.add(EventAttendance(
            batch_id=batch.id,
            employee_id=emp.id if emp else None,
            badge_employee_id=badge,
            category=cat,
            event_date=event_date,
            punched_at=row['punched_at'],
            timestamp_ms=ts,
            latitude=row['latitude'],
            longitude=row['longitude'],
            barangay=loc.get('barangay'),
            municipality=loc.get('municipality'),
            province=loc.get('province'),
            location_text=loc.get('location_text'),
        ))
        inserted += 1

    batch.inserted_count = inserted
    batch.skipped_count = skipped
    batch.unmatched_count = unmatched
    db.session.commit()
    return batch
